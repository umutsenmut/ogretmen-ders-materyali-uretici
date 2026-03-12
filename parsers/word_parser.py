from docx import Document


class WordParser:
    COLUMN_KEYS = [
        "week_number",
        "unit_topic",
        "learning_outcomes",
        "outcome_description",
        "teaching_techniques",
        "tools_equipment",
        "assessment",
        "special_days",
    ]

    def parse(self, filepath):
        entries = []
        try:
            doc = Document(filepath)
            for table in doc.tables:
                start_row = 0
                if table.rows:
                    first_row_text = " ".join(
                        cell.text.strip().lower()
                        for cell in table.rows[0].cells
                    )
                    if any(
                        kw in first_row_text
                        for kw in ["hafta", "week", "ünite", "konu"]
                    ):
                        start_row = 1
                for row in table.rows[start_row:]:
                    cells = [cell.text.strip() for cell in row.cells]
                    if all(c == "" for c in cells):
                        continue
                    entry = self._cells_to_entry(cells)
                    if entry:
                        entries.append(entry)
            # If no tables found, parse paragraphs
            if not entries:
                for para in doc.paragraphs:
                    text = para.text.strip()
                    if text:
                        entries.append(
                            {
                                "week_number": "",
                                "unit_topic": text,
                                "learning_outcomes": "",
                                "outcome_description": "",
                                "teaching_techniques": "",
                                "tools_equipment": "",
                                "assessment": "",
                                "special_days": "",
                            }
                        )
        except Exception as e:
            entries.append(
                {
                    "week_number": "Hata",
                    "unit_topic": f"Dosya okunamadı: {str(e)}",
                    "learning_outcomes": "",
                    "outcome_description": "",
                    "teaching_techniques": "",
                    "tools_equipment": "",
                    "assessment": "",
                    "special_days": "",
                }
            )
        return entries

    def _cells_to_entry(self, cells):
        while len(cells) < 8:
            cells.append("")
        entry = {}
        for i, key in enumerate(self.COLUMN_KEYS):
            entry[key] = cells[i] if i < len(cells) else ""
        if all(v == "" for v in entry.values()):
            return None
        return entry
