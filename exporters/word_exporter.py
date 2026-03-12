import io
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_heading_color(run, hex_color="0d6efd"):
    r, g, b = int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
    run.font.color.rgb = RGBColor(r, g, b)


def _add_heading(doc, text, level=1, color="0d6efd"):
    heading = doc.add_heading(level=level)
    run = heading.add_run(text)
    _set_heading_color(run, color)
    run.font.bold = True
    return heading


def _add_paragraph(doc, text, bold=False, size=10):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    return para


class WordExporter:
    def export_test(self, test_data, subject, topic):
        doc = Document()

        # Page margins
        for section in doc.sections:
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin = Cm(2.5)
            section.right_margin = Cm(2.5)

        _add_heading(doc, f"Test / Sınav", level=1)
        _add_heading(doc, f"{subject} - {topic}", level=2, color="198754")

        doc.add_paragraph(
            "Ad Soyad: ___________________________   Tarih: __________   Sınıf: _______"
        )
        doc.add_paragraph()

        multiple_choice = test_data.get("multiple_choice", [])
        open_ended = test_data.get("open_ended", [])

        if multiple_choice:
            _add_heading(doc, "A Bölümü - Çoktan Seçmeli Sorular", level=2)
            for i, q in enumerate(multiple_choice, 1):
                _add_paragraph(doc, f"{i}. {q.get('question', '')}", bold=False, size=10)
                opts = q.get("options", {})
                for opt_key in ["A", "B", "C", "D"]:
                    _add_paragraph(
                        doc, f"    {opt_key}) {opts.get(opt_key, '')}", size=10
                    )
                doc.add_paragraph()

        if open_ended:
            doc.add_page_break()
            _add_heading(doc, "B Bölümü - Açık Uçlu Sorular", level=2)
            for i, q in enumerate(open_ended, 1):
                _add_paragraph(doc, f"{i}. {q.get('question', '')}", bold=True, size=10)
                for _ in range(4):
                    doc.add_paragraph("_" * 70)
                doc.add_paragraph()

        # Answer key
        if multiple_choice:
            doc.add_page_break()
            _add_heading(doc, "Cevap Anahtarı", level=2)
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            for j, label in enumerate(["Soru", "Cevap", "Soru", "Cevap", "Soru"]):
                hdr[j].text = label
                hdr[j].paragraphs[0].runs[0].bold = True
            row_data = []
            for i, q in enumerate(multiple_choice, 1):
                row_data.append((str(i), q.get("correct_answer", "-")))
            # Fill table rows
            for k in range(0, len(row_data), 2):
                cells = table.add_row().cells
                cells[0].text = row_data[k][0]
                cells[1].text = row_data[k][1]
                if k + 1 < len(row_data):
                    cells[2].text = row_data[k + 1][0]
                    cells[3].text = row_data[k + 1][1]
                else:
                    cells[2].text = ""
                    cells[3].text = ""
                cells[4].text = ""

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    def export_notes(self, notes_data, subject, topic):
        doc = Document()

        for section in doc.sections:
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin = Cm(2.5)
            section.right_margin = Cm(2.5)

        doc_title = notes_data.get("title", f"{subject} - {topic} Öğretmen Notları")
        _add_heading(doc, doc_title, level=1)

        main_concepts = notes_data.get("main_concepts", [])
        if main_concepts:
            _add_heading(doc, "Ana Kavramlar", level=2)
            for concept in main_concepts:
                _add_paragraph(doc, f"• {concept}", size=10)
            doc.add_paragraph()

        detailed = notes_data.get("detailed_explanation", "")
        if detailed:
            _add_heading(doc, "Ayrıntılı Açıklama", level=2)
            para = doc.add_paragraph(detailed)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            doc.add_paragraph()

        examples = notes_data.get("examples", [])
        if examples:
            _add_heading(doc, "Örnekler", level=2)
            for ex in examples:
                _add_paragraph(doc, f"• {ex}", size=10)
            doc.add_paragraph()

        warnings = notes_data.get("warnings", [])
        if warnings:
            _add_heading(doc, "Dikkat Edilmesi Gerekenler", level=2, color="dc3545")
            for w in warnings:
                _add_paragraph(doc, f"⚠ {w}", size=10)
            doc.add_paragraph()

        tips = notes_data.get("tips", [])
        if tips:
            _add_heading(doc, "Öğretim İpuçları", level=2, color="198754")
            for tip in tips:
                _add_paragraph(doc, f"✓ {tip}", size=10)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
